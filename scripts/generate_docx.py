# -*- coding: utf-8 -*-
"""
generate_docx.py
Converts manuscript/master.md into publication-quality DOCX with RTL support,
Word styles, callouts, tables, headers, and footers.
"""

import re
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def sanitize_text(s):
    if not s:
        return ""
    # strip XML invalid control characters
    return "".join(ch for ch in str(s) if ord(ch) >= 32 or ch in "\n\r\t")

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for b_name, b_val in borders.items():
        if b_val:
            b_elem = parse_xml(f'<w:{b_name} {nsdecls("w")} w:val="{b_val.get("val", "single")}" '
                               f'w:sz="{b_val.get("sz", "4")}" w:space="0" '
                               f'w:color="{b_val.get("color", "D3D3D3")}"/>')
            tcBorders.append(b_elem)
        else:
            b_elem = parse_xml(f'<w:{b_name} {nsdecls("w")} w:val="none"/>')
            tcBorders.append(b_elem)
    tcPr.append(tcBorders)

def add_styled_paragraph(doc, text, style_type="body", space_before=2, space_after=4):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pPr.append(OxmlElement('w:bidi'))
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15

    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\$\$.*?\$\$)', text)
    for token in tokens:
        if not token:
            continue
        is_bold = False
        is_italic = False
        is_code = False
        is_math = False
        clean_text = token
        if token.startswith('**') and token.endswith('**'):
            is_bold = True
            clean_text = token[2:-2]
        elif token.startswith('*') and token.endswith('*'):
            is_italic = True
            clean_text = token[1:-1]
        elif token.startswith('`') and token.endswith('`'):
            is_code = True
            clean_text = token[1:-1]
        elif token.startswith('$$') and token.endswith('$$'):
            is_math = True
            clean_text = token[2:-2]

        clean_text = sanitize_text(clean_text)
        run = p.add_run(clean_text)
        run.font.name = "Vazirmatn"
        rPr = run._r.get_or_add_rPr()
        rPr.append(OxmlElement('w:rtl'))
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="Calibri" w:hAnsi="Calibri" w:cs="Vazirmatn"/>')
        rPr.append(rFonts)

        if style_type == "h1":
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 44, 89)
        elif style_type == "h2":
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 85, 140)
        elif style_type == "h3":
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 90, 130)
        elif style_type == "h4":
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(50, 50, 50)
        elif style_type == "bullet":
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor(40, 40, 40)
        else:
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor(35, 35, 35)

        if is_bold:
            run.font.bold = True
        if is_italic:
            run.font.italic = True
        if is_code or is_math:
            run.font.name = "Courier New"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(160, 40, 40)
            
    return p

def add_callout_box(doc, text, callout_type="info"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl._tbl.tblPr.append(OxmlElement('w:bidiVisual'))
    cell = tbl.cell(0, 0)
    
    border_color = "005A9C"
    bg_color = "F0F5FA"
    if "خطر" in text or "هشدار" in text or "Red Lines" in text or "خطوط سرخ" in text:
        border_color = "C0392B"
        bg_color = "FDF2E9"
    elif "طلایی" in text or "Memory Hook" in text or "فرمول" in text or "کلید" in text:
        border_color = "D4AC0D"
        bg_color = "FEF9E7"
    elif "پاسخ نمونه" in text or "Model Answer" in text:
        border_color = "27AE60"
        bg_color = "EAFAF1"
        
    set_cell_shading(cell, bg_color)
    set_cell_borders(cell, 
                     top={'val': 'single', 'sz': '4', 'color': 'E0E0E0'},
                     bottom={'val': 'single', 'sz': '4', 'color': 'E0E0E0'},
                     left={'val': 'single', 'sz': '4', 'color': 'E0E0E0'},
                     right={'val': 'single', 'sz': '24', 'color': border_color})
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    p = cell.paragraphs[0]
    pPr = p._p.get_or_add_pPr()
    pPr.append(OxmlElement('w:bidi'))
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.line_spacing = 1.15
    
    clean_lines = [l.strip().lstrip('>').strip() for l in text.split('\n') if l.strip()]
    for i, line in enumerate(clean_lines):
        if i > 0:
            p = cell.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pPr.append(OxmlElement('w:bidi'))
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.line_spacing = 1.15
            
        tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\$\$.*?\$\$)', line)
        for token in tokens:
            if not token:
                continue
            is_bold = token.startswith('**') and token.endswith('**')
            is_code = token.startswith('`') and token.endswith('`')
            is_math = token.startswith('$$') and token.endswith('$$')
            ctext = token[2:-2] if (is_bold or is_math) else (token[1:-1] if is_code else token)
            ctext = sanitize_text(ctext)
            run = p.add_run(ctext)
            run.font.name = "Vazirmatn"
            run.font.size = Pt(10)
            if is_bold:
                run.font.bold = True
                run.font.color.rgb = RGBColor(20, 20, 20)
            elif is_math:
                run.font.bold = True
                run.font.color.rgb = RGBColor(140, 20, 20)
            else:
                run.font.color.rgb = RGBColor(40, 40, 40)
            rPr = run._r.get_or_add_rPr()
            rPr.append(OxmlElement('w:rtl'))
            rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="Calibri" w:hAnsi="Calibri" w:cs="Vazirmatn"/>')
            rPr.append(rFonts)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(2)

def add_markdown_table(doc, table_lines):
    rows_data = []
    for line in table_lines:
        line = line.strip()
        if not line or set(line).issubset({'-', '|', ':', ' '}):
            continue
        cells = [c.strip() for c in line.strip('|').split('|')]
        rows_data.append(cells)
        
    if not rows_data:
        return
        
    col_count = max(len(r) for r in rows_data)
    table = doc.add_table(rows=len(rows_data), cols=col_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table._tbl.tblPr.append(OxmlElement('w:bidiVisual'))
    
    for r_idx, row in enumerate(rows_data):
        is_header = (r_idx == 0)
        for c_idx in range(col_count):
            cell = table.cell(r_idx, c_idx)
            val = row[c_idx] if c_idx < len(row) else ""
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            
            if is_header:
                set_cell_shading(cell, "1B365D")
                text_color = RGBColor(255, 255, 255)
                set_cell_borders(cell, 
                                 top={'val': 'single', 'sz': '4', 'color': '1B365D'},
                                 bottom={'val': 'single', 'sz': '12', 'color': '002244'},
                                 left={'val': 'single', 'sz': '4', 'color': '2E4A70'},
                                 right={'val': 'single', 'sz': '4', 'color': '2E4A70'})
            else:
                bg = "F9FAFC" if r_idx % 2 == 1 else "FFFFFF"
                set_cell_shading(cell, bg)
                text_color = RGBColor(30, 30, 30)
                set_cell_borders(cell, 
                                 top={'val': 'single', 'sz': '4', 'color': 'E0E0E0'},
                                 bottom={'val': 'single', 'sz': '4', 'color': 'E0E0E0'},
                                 left={'val': 'single', 'sz': '4', 'color': 'E0E0E0'},
                                 right={'val': 'single', 'sz': '4', 'color': 'E0E0E0'})
                
            p = cell.paragraphs[0]
            pPr = p._p.get_or_add_pPr()
            pPr.append(OxmlElement('w:bidi'))
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if not is_header else WD_ALIGN_PARAGRAPH.CENTER
            
            sub_lines = val.replace('<br>', '\n').split('\n')
            for s_idx, s_line in enumerate(sub_lines):
                if s_idx > 0:
                    p = cell.add_paragraph()
                    pPr = p._p.get_or_add_pPr()
                    pPr.append(OxmlElement('w:bidi'))
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*)', s_line)
                for token in tokens:
                    if not token:
                        continue
                    is_bold = token.startswith('**') and token.endswith('**')
                    ctext = token[2:-2] if is_bold else token
                    ctext = sanitize_text(ctext)
                    run = p.add_run(ctext)
                    run.font.name = "Vazirmatn"
                    run.font.size = Pt(9.5 if not is_header else 10)
                    run.font.bold = is_bold or is_header
                    run.font.color.rgb = text_color
                    rPr = run._r.get_or_add_rPr()
                    rPr.append(OxmlElement('w:rtl'))
                    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="Calibri" w:hAnsi="Calibri" w:cs="Vazirmatn"/>')
                    rPr.append(rFonts)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(2)

def build_docx_from_markdown(md_path, docx_path):
    print(f"Reading {md_path}...")
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()
        
    doc = docx.Document()
    
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hrun = hp.add_run("Shuhada Organization — Provincial Coordinator (Daikundi) | 24-Hour Master Guide")
    hrun.font.name = "Vazirmatn"
    hrun.font.size = Pt(8.5)
    hrun.font.color.rgb = RGBColor(120, 120, 120)
    
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    frun = fp.add_run("صفحه — ")
    frun.font.name = "Vazirmatn"
    frun.font.size = Pt(9)
    frun.font.color.rgb = RGBColor(120, 120, 120)
    rPr = frun._r.get_or_add_rPr()
    rPr.append(OxmlElement('w:rtl'))
    
    fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
    fp._p.append(fldSimple)

    lines = content.split('\n')
    i = 0
    in_front_matter = False
    in_code_block = False
    code_block_lines = []
    table_lines = []
    quote_lines = []
    
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        
        if i == 0 and stripped == '---':
            in_front_matter = True
            i += 1
            continue
        if in_front_matter:
            if stripped == '---':
                in_front_matter = False
            i += 1
            continue
            
        if stripped.startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_block_lines = []
            else:
                in_code_block = False
                add_callout_box(doc, "\n".join(code_block_lines), callout_type="code")
            i += 1
            continue
            
        if in_code_block:
            code_block_lines.append(line)
            i += 1
            continue
            
        if stripped.startswith('|') and stripped.endswith('|'):
            table_lines.append(stripped)
            i += 1
            continue
        else:
            if table_lines:
                add_markdown_table(doc, table_lines)
                table_lines = []
                
        if stripped.startswith('>'):
            quote_lines.append(stripped)
            i += 1
            continue
        else:
            if quote_lines:
                add_callout_box(doc, "\n".join(quote_lines))
                quote_lines = []
                
        if not stripped:
            i += 1
            continue
            
        if stripped.startswith('# '):
            text = stripped[2:].strip()
            add_styled_paragraph(doc, text, style_type="h1", space_before=14, space_after=6)
        elif stripped.startswith('## '):
            text = stripped[3:].strip()
            add_styled_paragraph(doc, text, style_type="h2", space_before=10, space_after=4)
        elif stripped.startswith('### '):
            text = stripped[4:].strip()
            add_styled_paragraph(doc, text, style_type="h3", space_before=8, space_after=3)
        elif stripped.startswith('#### '):
            text = stripped[5:].strip()
            add_styled_paragraph(doc, text, style_type="h4", space_before=6, space_after=2)
        elif stripped.startswith('* ') or stripped.startswith('- '):
            text = "• " + stripped[2:].strip()
            add_styled_paragraph(doc, text, style_type="bullet", space_before=1, space_after=2)
        elif re.match(r'^\d+\.\s', stripped):
            add_styled_paragraph(doc, stripped, style_type="body", space_before=1, space_after=2)
        elif stripped.startswith('---'):
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="CCCCCC"/></w:pBdr>')
            pPr.append(pBdr)
        else:
            add_styled_paragraph(doc, stripped, style_type="body", space_before=2, space_after=3)
            
        i += 1
        
    if table_lines:
        add_markdown_table(doc, table_lines)
    if quote_lines:
        add_callout_box(doc, "\n".join(quote_lines))
        
    os.makedirs(os.path.dirname(docx_path), exist_ok=True)
    doc.save(docx_path)
    print(f"Generated DOCX saved successfully at: {docx_path}")
    print(f"File size: {os.path.getsize(docx_path):,} bytes")

if __name__ == "__main__":
    md_in = "manuscript/master.md"
    docx_out = "build/Provincial_Coordinator_24Hour_Exam_Master_Guide.docx"
    build_docx_from_markdown(md_in, docx_out)
