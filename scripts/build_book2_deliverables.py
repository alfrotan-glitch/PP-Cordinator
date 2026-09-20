# -*- coding: utf-8 -*-
"""
build_book2_deliverables.py
Compiles manuscript/health_management_master.md into:
1. build/Health_Management_Master_Guide.docx
2. کتاب_Health_Management_نسخه_ویرایش‌شده.docx
3. build/health_management.html
4. build/Health_Management_Master_Guide.epub
"""

import os
import sys
import shutil
import re
import html
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from ebooklib import epub

def sanitize_text(s):
    if not s:
        return ""
    return "".join(ch for ch in str(s) if ord(ch) >= 32 or ch in "\n\r\t")

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for b_name, b_val in borders.items():
        if b_val:
            b_elem = parse_xml(f'<w:{b_name} {nsdecls("w")} w:val="{b_val.get("val", "single")}" '
                               f'w:sz="{b_val.get("sz", "4")}" w:space="0" '
                               f'w:color="{b_val.get("color", "D3D3D3")}"/>')
            tcBorders.append(b_elem)
        else:
            b_elem = parse_xml(f'<w:{b_name} {nsdecls("w")} w:val="none"/>')
            tcBorders.append(b_elem)
    tcPr.append(tcBorders)

def add_styled_paragraph(doc, text, style_type="body", space_before=2, space_after=4):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pPr.append(OxmlElement('w:bidi'))
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15

    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\$\$.*?\$\$)', text)
    for token in tokens:
        if not token:
            continue
        is_bold = False
        is_italic = False
        is_code = False
        is_math = False
        clean_text = token
        if token.startswith('**') and token.endswith('**'):
            is_bold = True
            clean_text = token[2:-2]
        elif token.startswith('*') and token.endswith('*'):
            is_italic = True
            clean_text = token[1:-1]
        elif token.startswith('`') and token.endswith('`'):
            is_code = True
            clean_text = token[1:-1]
        elif token.startswith('$$') and token.endswith('$$'):
            is_math = True
            clean_text = token[2:-2]

        clean_text = sanitize_text(clean_text)
        run = p.add_run(clean_text)
        run.font.name = "Vazirmatn"
        rPr = run._r.get_or_add_rPr()
        rPr.append(OxmlElement('w:rtl'))
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="Calibri" w:hAnsi="Calibri" w:cs="Vazirmatn"/>')
        rPr.append(rFonts)

        if style_type == "h1":
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 44, 89)
        elif style_type == "h2":
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 85, 140)
        elif style_type == "h3":
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 90, 130)
        elif style_type == "h4":
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(50, 50, 50)
        elif style_type == "bullet":
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor(40, 40, 40)
        else:
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(30, 30, 30)

        if is_bold:
            run.font.bold = True
        if is_italic:
            run.font.italic = True
        if is_code or is_math:
            run.font.name = "Consolas"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(120, 20, 20)

def add_callout_box(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)

    cell = table.cell(0, 0)
    set_cell_shading(cell, "F0F4F8")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=220)
    
    border_right = {'val': 'single', 'sz': '24', 'color': '0F2C59'}
    set_cell_borders(cell, top=None, bottom=None, left=None, right=border_right)

    p = cell.paragraphs[0]
    pPr = p._p.get_or_add_pPr()
    pPr.append(OxmlElement('w:bidi'))
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)

    run = p.add_run(sanitize_text(text.strip()))
    run.font.name = "Vazirmatn"
    run.font.size = Pt(10)
    run.font.italic = True
    run.font.color.rgb = RGBColor(27, 54, 93)
    rPr = run._r.get_or_add_rPr()
    rPr.append(OxmlElement('w:rtl'))

def add_markdown_table(doc, table_lines):
    rows_data = []
    for line in table_lines:
        line_clean = line.strip().strip('|')
        if not line_clean or re.match(r'^[\s\-\:\.\=\|]+$', line_clean):
            continue
        cols = [c.strip() for c in line.split('|')[1:-1]]
        if cols:
            rows_data.append(cols)

    if not rows_data:
        return

    num_cols = max(len(r) for r in rows_data)
    table = doc.add_table(rows=len(rows_data), cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for col_idx in range(num_cols):
        table.columns[col_idx].width = Inches(6.5 / num_cols)

    for r_idx, row in enumerate(rows_data):
        is_header = (r_idx == 0)
        for c_idx in range(num_cols):
            cell = table.cell(r_idx, c_idx)
            val = row[c_idx] if c_idx < len(row) else ""
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            bdr = {'val': 'single', 'sz': '4', 'color': 'CCCCCC'}
            set_cell_borders(cell, top=bdr, bottom=bdr, left=bdr, right=bdr)

            if is_header:
                set_cell_shading(cell, "0F2C59")
            else:
                if r_idx % 2 == 1:
                    set_cell_shading(cell, "FFFFFF")
                else:
                    set_cell_shading(cell, "F9FBFD")

            p = cell.paragraphs[0]
            pPr = p._p.get_or_add_pPr()
            pPr.append(OxmlElement('w:bidi'))
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.line_spacing = 1.05
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)

            run = p.add_run(sanitize_text(val))
            run.font.name = "Vazirmatn"
            run.font.size = Pt(9.5)
            rPr = run._r.get_or_add_rPr()
            rPr.append(OxmlElement('w:rtl'))

            if is_header:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                run.font.color.rgb = RGBColor(30, 30, 30)

def generate_docx(md_path, docx_path):
    print(f"Building DOCX: {docx_path}...")
    doc = docx.Document()
    
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
        sectPr = section._sectPr
        bidi = OxmlElement('w:bidi')
        sectPr.append(bidi)
        
        header = section.header
        hp = header.paragraphs[0]
        hpPr = hp._p.get_or_add_pPr()
        hpPr.append(OxmlElement('w:bidi'))
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        hrun = hp.add_run("مدیریت صحی — رهنما و مرجع جامع ۲۴ ساعته آزمون‌های استخدامی و کار عملی در افغانستان")
        hrun.font.name = "Vazirmatn"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(120, 120, 120)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fpPr = fp._p.get_or_add_pPr()
        fpPr.append(OxmlElement('w:bidi'))
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Health Management — 24-Hour Exam & Field Practice Master Guide | سازمان شهدا")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(120, 120, 120)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_lines = []
    table_lines = []
    quote_lines = []
    in_front_matter = False

    i = 0
    n = len(lines)
    while i < n:
        raw_line = lines[i]
        line = raw_line.rstrip('\r\n')
        stripped = line.strip()

        if i == 0 and stripped == '---':
            in_front_matter = True
            i += 1
            continue
        if in_front_matter:
            if stripped == '---':
                in_front_matter = False
            i += 1
            continue

        if stripped.startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_lines = []
            else:
                in_code_block = False
                callout_text = "\n".join(code_lines)
                add_callout_box(doc, callout_text)
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        if stripped.startswith('|') and stripped.endswith('|'):
            table_lines.append(stripped)
            i += 1
            continue
        else:
            if table_lines:
                add_markdown_table(doc, table_lines)
                table_lines = []

        if stripped.startswith('>'):
            quote_text = stripped.lstrip('>').strip()
            quote_lines.append(quote_text)
            i += 1
            continue
        else:
            if quote_lines:
                add_callout_box(doc, "\n".join(quote_lines))
                quote_lines = []

        if not stripped:
            i += 1
            continue

        if stripped.startswith('# '):
            text = stripped[2:].strip()
            add_styled_paragraph(doc, text, style_type="h1", space_before=14, space_after=6)
        elif stripped.startswith('## '):
            text = stripped[3:].strip()
            add_styled_paragraph(doc, text, style_type="h2", space_before=10, space_after=4)
        elif stripped.startswith('### '):
            text = stripped[4:].strip()
            add_styled_paragraph(doc, text, style_type="h3", space_before=8, space_after=3)
        elif stripped.startswith('#### '):
            text = stripped[5:].strip()
            add_styled_paragraph(doc, text, style_type="h4", space_before=6, space_after=2)
        elif stripped.startswith('* ') or stripped.startswith('- '):
            text = "• " + stripped[2:].strip()
            add_styled_paragraph(doc, text, style_type="bullet", space_before=1, space_after=2)
        elif re.match(r'^\d+\.\s', stripped):
            add_styled_paragraph(doc, stripped, style_type="body", space_before=1, space_after=2)
        elif stripped.startswith('---'):
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="CCCCCC"/></w:pBdr>')
            pPr.append(pBdr)
        else:
            add_styled_paragraph(doc, stripped, style_type="body", space_before=2, space_after=3)

        i += 1

    if table_lines:
        add_markdown_table(doc, table_lines)
    if quote_lines:
        add_callout_box(doc, "\n".join(quote_lines))

    os.makedirs(os.path.dirname(docx_path), exist_ok=True)
    doc.save(docx_path)
    print(f"Generated DOCX: {docx_path} ({os.path.getsize(docx_path):,} bytes)")

def generate_html(md_path, html_path):
    print(f"Building HTML: {html_path}...")
    with open(md_path, 'r', encoding='utf-8') as f:
        md = f.read()

    lines = md.split('\n')
    body_html = []
    
    in_code = False
    code_lines = []
    table_lines = []
    quote_lines = []
    in_front_matter = False
    
    def process_inline(text):
        text = html.escape(text)
        text = re.sub(r'\*\*(.*?)\*\*', r'<strong>\1</strong>', text)
        text = re.sub(r'\*(.*?)\*', r'<em>\1</em>', text)
        text = re.sub(r'`(.*?)`', r'<code>\1</code>', text)
        text = re.sub(r'\$\$(.*?)\$\$', r'<span class="math-formula">\1</span>', text)
        return text

    for i, line in enumerate(lines):
        stripped = line.strip()
        
        if i == 0 and stripped == '---':
            in_front_matter = True
            continue
        if in_front_matter:
            if stripped == '---':
                in_front_matter = False
            continue

        if stripped.startswith('```'):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                in_code = False
                block_content = html.escape("\n".join(code_lines))
                body_html.append(f'<pre><code>{block_content}</code></pre>')
                code_lines = []
            continue

        if in_code:
            code_lines.append(line)
            continue

        if stripped.startswith('|') and stripped.endswith('|'):
            table_lines.append(stripped)
            continue
        else:
            if table_lines:
                table_html = ['<div class="table-container"><table>']
                rows = [r.strip().strip('|') for r in table_lines if r.strip().strip('|')]
                header_done = False
                for row_idx, r in enumerate(rows):
                    if re.match(r'^[\s\-\:\.\=\|]+$', r):
                        continue
                    cols = [c.strip() for c in r.split('|')]
                    if not header_done:
                        table_html.append('<thead><tr>')
                        for c in cols:
                            table_html.append(f'<th>{process_inline(c)}</th>')
                        table_html.append('</tr></thead><tbody>')
                        header_done = True
                    else:
                        table_html.append('<tr>')
                        for c in cols:
                            table_html.append(f'<td>{process_inline(c)}</td>')
                        table_html.append('</tr>')
                if header_done:
                    table_html.append('</tbody>')
                table_html.append('</table></div>')
                body_html.append('\n'.join(table_html))
                table_lines = []

        if stripped.startswith('>'):
            quote_text = stripped.lstrip('>').strip()
            quote_lines.append(quote_text)
            continue
        else:
            if quote_lines:
                q_content = '<br>'.join([process_inline(q) for q in quote_lines])
                body_html.append(f'<blockquote class="callout">{q_content}</blockquote>')
                quote_lines = []

        if not stripped:
            continue

        if stripped.startswith('# '):
            title = process_inline(stripped[2:])
            body_html.append(f'<h1 class="part-title">{title}</h1>')
        elif stripped.startswith('## '):
            title = process_inline(stripped[3:])
            body_html.append(f'<h2 class="chapter-title">{title}</h2>')
        elif stripped.startswith('### '):
            title = process_inline(stripped[4:])
            body_html.append(f'<h3 class="section-title">{title}</h3>')
        elif stripped.startswith('#### '):
            title = process_inline(stripped[5:])
            body_html.append(f'<h4 class="sub-section-title">{title}</h4>')
        elif stripped.startswith('* ') or stripped.startswith('- '):
            item = process_inline(stripped[2:])
            body_html.append(f'<li class="bullet-item">{item}</li>')
        elif re.match(r'^\d+\.\s', stripped):
            item = process_inline(re.sub(r'^\d+\.\s*', '', stripped))
            body_html.append(f'<li class="numbered-item">{item}</li>')
        elif stripped.startswith('---'):
            body_html.append('<hr class="divider">')
        else:
            p_text = process_inline(stripped)
            body_html.append(f'<p>{p_text}</p>')

    content = '\n'.join(body_html)

    html_template = f"""<!DOCTYPE html>
<html lang="fa" dir="rtl">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Health Management — 24-Hour Exam & Field Practice Master Guide</title>
    <link rel="preconnect" href="https://fonts.googleapis.com">
    <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
    <link href="https://fonts.googleapis.com/css2?family=Vazirmatn:wght@300;400;600;700;800;900&display=swap" rel="stylesheet">
    <style>
        :root {{
            --primary: #0F2C59;
            --primary-light: #005A9C;
            --accent: #C28B14;
            --bg-color: #F8F9FA;
            --card-bg: #FFFFFF;
            --text-main: #1F2937;
            --text-muted: #4B5563;
            --border-color: #E5E7EB;
        }}
        * {{ box-sizing: border-box; }}
        body {{
            font-family: 'Vazirmatn', -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif;
            background-color: var(--bg-color);
            color: var(--text-main);
            margin: 0;
            padding: 0;
            line-height: 1.8;
            font-size: 16px;
            direction: rtl;
            text-align: right;
        }}
        .header {{
            background: linear-gradient(135deg, #0F2C59 0%, #173B6C 100%);
            color: #FFFFFF;
            padding: 40px 20px;
            text-align: center;
            border-bottom: 5px solid var(--accent);
            box-shadow: 0 4px 15px rgba(0,0,0,0.1);
        }}
        .header h1 {{ margin: 0 0 10px 0; font-size: 28px; font-weight: 900; }}
        .header p {{ margin: 5px 0 0 0; font-size: 16px; opacity: 0.9; }}
        .container {{
            max-width: 960px;
            margin: 30px auto;
            background: var(--card-bg);
            padding: 45px 55px;
            border-radius: 12px;
            box-shadow: 0 10px 30px rgba(0,0,0,0.05);
            border: 1px solid var(--border-color);
        }}
        .search-box {{
            width: 100%;
            padding: 12px 18px;
            font-size: 15px;
            border: 2px solid var(--border-color);
            border-radius: 8px;
            margin-bottom: 30px;
            direction: rtl;
            font-family: inherit;
        }}
        .search-box:focus {{ border-color: var(--primary-light); outline: none; }}
        h1.part-title {{
            color: var(--primary);
            font-size: 24px;
            font-weight: 800;
            border-right: 6px solid var(--accent);
            padding-right: 15px;
            margin-top: 50px;
            margin-bottom: 25px;
            background: #F0F4F8;
            padding: 12px 18px;
            border-radius: 4px;
        }}
        h2.chapter-title {{
            color: var(--primary-light);
            font-size: 20px;
            font-weight: 700;
            margin-top: 35px;
            margin-bottom: 15px;
            border-bottom: 2px solid var(--border-color);
            padding-bottom: 8px;
        }}
        h3.section-title {{ color: #1E3A8A; font-size: 17px; font-weight: 600; margin-top: 25px; }}
        h4.sub-section-title {{ color: #374151; font-size: 15px; font-weight: 600; margin-top: 15px; }}
        p {{ margin-bottom: 16px; color: var(--text-main); }}
        .bullet-item, .numbered-item {{ margin-bottom: 8px; color: var(--text-main); margin-right: 25px; }}
        .callout {{
            border-right: 4px solid var(--primary-light);
            background-color: #F0F4F8;
            padding: 16px 20px;
            margin: 25px 0;
            border-radius: 0 8px 8px 0;
            font-style: italic;
            color: #1E293B;
        }}
        .table-container {{ overflow-x: auto; margin: 25px 0; }}
        table {{
            width: 100%;
            border-collapse: collapse;
            font-size: 14.5px;
            background: #FFFFFF;
        }}
        th, td {{
            padding: 10px 14px;
            border: 1px solid var(--border-color);
            text-align: right;
        }}
        th {{
            background-color: var(--primary);
            color: #FFFFFF;
            font-weight: 600;
        }}
        tr:nth-child(even) {{ background-color: #F9FAFB; }}
        pre {{
            background-color: #1E293B;
            color: #F8FAFC;
            padding: 16px 20px;
            border-radius: 8px;
            overflow-x: auto;
            direction: ltr;
            text-align: left;
            font-size: 13.5px;
        }}
        code {{
            background: #EEF2F6;
            color: #B91C1C;
            padding: 2px 6px;
            border-radius: 4px;
            font-family: monospace;
            font-size: 14px;
        }}
        .math-formula {{
            background: #EFF6FF;
            color: #1D4ED8;
            padding: 3px 8px;
            border-radius: 4px;
            font-family: monospace;
            font-weight: 600;
        }}
        .divider {{
            border: none;
            border-top: 1px solid var(--border-color);
            margin: 40px 0;
        }}
        .footer {{
            text-align: center;
            padding: 30px;
            color: var(--text-muted);
            font-size: 14px;
        }}
    </style>
</head>
<body>
    <div class="header">
        <h1>مدیریت صحی — رهنما و مرجع جامع ۲۴ ساعته آزمون‌های استخدامی و کار عملی در افغانستان</h1>
        <p>HEALTH MANAGEMENT — 24-HOUR EXAM & FIELD PRACTICE MASTER GUIDE</p>
        <p>سازمان شهدا (Shuhada Organization) — ولایت دایکندی (نیلی)</p>
    </div>
    <div class="container">
        <input type="text" id="searchBox" class="search-box" placeholder="جستجوی موضوعات، کلمات کلیدی یا اصطلاحات صحی..." onkeyup="filterContent()">
        <div id="contentBody">
            {content}
        </div>
    </div>
    <div class="footer">
        Health Management Master Guide &copy; 2026 | سازمان شهدا — تدوین شده بر اساس آخرین رهنمودهای وزارت صحت عامه
    </div>
    <script>
        function filterContent() {{
            var input = document.getElementById('searchBox');
            var filter = input.value.toLowerCase();
            var body = document.getElementById('contentBody');
            var paragraphs = body.getElementsByTagName('p');
            for (var i = 0; i < paragraphs.length; i++) {{
                var text = paragraphs[i].textContent || paragraphs[i].innerText;
                if (text.toLowerCase().indexOf(filter) > -1 || filter === '') {{
                    paragraphs[i].style.display = "";
                }} else {{
                    paragraphs[i].style.display = "none";
                }}
            }}
        }}
    </script>
</body>
</html>
"""
    os.makedirs(os.path.dirname(html_path), exist_ok=True)
    with open(html_path, 'w', encoding='utf-8') as f:
        f.write(html_template)
    print(f"Generated HTML: {html_path} ({os.path.getsize(html_path):,} bytes)")

def generate_epub(md_path, epub_path):
    print(f"Building EPUB: {epub_path}...")
    with open(md_path, 'r', encoding='utf-8') as f:
        md = f.read()

    book = epub.EpubBook()
    book.set_identifier('shuhada-health-management-master-2026')
    book.set_title('HEALTH MANAGEMENT — 24-HOUR EXAM & FIELD PRACTICE MASTER GUIDE')
    book.set_language('fa')
    book.add_author('Shuhada Organization — Health Directorate')
    
    style = '''
    @namespace epub "http://www.idpf.org/2007/ops";
    body {
        font-family: sans-serif;
        direction: rtl;
        text-align: right;
        line-height: 1.6;
        padding: 5%;
    }
    h1 { color: #0F2C59; border-bottom: 2px solid #005A9C; }
    h2 { color: #005A9C; }
    h3 { color: #1E3A8A; }
    .callout {
        border-right: 4px solid #005A9C;
        background-color: #F0F4F8;
        padding: 10px;
        margin: 15px 0;
    }
    table { width: 100%; border-collapse: collapse; margin: 15px 0; }
    th, td { border: 1px solid #CCC; padding: 6px; text-align: right; }
    th { background-color: #0F2C59; color: white; }
    '''
    default_css = epub.EpubItem(uid="style_nav", file_name="style/nav.css", media_type="text/css", content=style)
    book.add_item(default_css)

    chapters = []
    parts_raw = re.split(r'\n# ', md)
    
    # Intro chapter
    intro_html = f"<html><head><link rel=\"stylesheet\" href=\"style/nav.css\"/></head><body><h1>مقدمه و راهنما</h1><p>{html.escape(parts_raw[0][:2000])}</p></body></html>"
    c_intro = epub.EpubHtml(title="Introduction", file_name="intro.xhtml", lang="fa")
    c_intro.content = intro_html
    c_intro.add_item(default_css)
    book.add_item(c_intro)
    chapters.append(c_intro)

    for idx, p in enumerate(parts_raw[1:], start=1):
        lines = p.split('\n')
        part_title = lines[0].strip()
        part_body = '\n'.join(lines[1:])
        
        # simple html conversion
        html_body = []
        for line in part_body.split('\n'):
            line_str = line.strip()
            if line_str.startswith('## '):
                html_body.append(f"<h2>{html.escape(line_str[3:])}</h2>")
            elif line_str.startswith('### '):
                html_body.append(f"<h3>{html.escape(line_str[4:])}</h3>")
            elif line_str.startswith('* ') or line_str.startswith('- '):
                html_body.append(f"<li>{html.escape(line_str[2:])}</li>")
            elif line_str:
                html_body.append(f"<p>{html.escape(line_str)}</p>")
                
        c_content = f"<html><head><link rel=\"stylesheet\" href=\"style/nav.css\"/></head><body><h1>{html.escape(part_title)}</h1>{''.join(html_body)}</body></html>"
        c = epub.EpubHtml(title=part_title[:40], file_name=f"part_{idx}.xhtml", lang="fa")
        c.content = c_content
        c.add_item(default_css)
        book.add_item(c)
        chapters.append(c)

    book.toc = tuple(chapters)
    book.spine = ['nav'] + chapters
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())

    os.makedirs(os.path.dirname(epub_path), exist_ok=True)
    epub.write_epub(epub_path, book, {})
    print(f"Generated EPUB: {epub_path} ({os.path.getsize(epub_path):,} bytes)")

def build_all():
    md_in = "manuscript/health_management_master.md"
    docx_build = "build/Health_Management_Master_Guide.docx"
    docx_persian = "کتاب_Health_Management_نسخه_ویرایش‌شده.docx"
    html_out = "build/health_management.html"
    epub_out = "build/Health_Management_Master_Guide.epub"
    
    print("=== STARTING DELIVERABLES BUILD FOR BOOK 2 ===")
    generate_docx(md_in, docx_build)
    
    shutil.copy(docx_build, docx_persian)
    print(f"Copied Persian-named DOCX: {docx_persian} ({os.path.getsize(docx_persian):,} bytes)")
    
    generate_html(md_in, html_out)
    generate_epub(md_in, epub_out)
    print("=== ALL DELIVERABLES BUILT SUCCESSFULLY ===")

if __name__ == "__main__":
    build_all()
